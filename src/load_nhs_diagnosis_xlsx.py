"""
load_nhs_diagnosis_xlsx.py

Parses the older-format NHS Digital diagnosis Excel files
(2021-22, 2022-23, 2023-24), which use a wide table layout with a header
row of named columns, unlike the 2024-25 long-format CSV.

Sheets used:
  - "Primary Diagnosis 3 Character" -> F50 total row
  - "Primary Diagnosis 4 Character" -> F50.0-F50.9 subtype rows
"""

import openpyxl
import pandas as pd

ED_CODE_LABELS = {
    "F50": "All Eating Disorders (total)",
    "F50.0": "Anorexia Nervosa",
    "F50.1": "Atypical Anorexia Nervosa",
    "F50.2": "Bulimia Nervosa",
    "F50.3": "Atypical Bulimia Nervosa",
    "F50.4": "Overeating (psychological)",
    "F50.5": "Vomiting (psychological)",
    "F50.8": "Other Eating Disorders (incl. Binge Eating Disorder)",
    "F50.9": "Eating Disorder, Unspecified",
}

FILES = {
    "2021-22": "data/raw/hosp-epis-stat-admi-diag-2021-22-tab.xlsx",
    "2022-23": "data/raw/hosp-epis-stat-admi-diag-2022-23-tab_V2.xlsx",
    "2023-24": "data/raw/hosp-epis-stat-admi-diag-2023-24-tab.xlsx",
}


def find_header_row(ws):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        if row and row[0] and "code and description" in str(row[0]):
            return i
    raise ValueError("Header row not found")


def parse_sheet(path, sheet_name, code_prefix_filter):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    header_row_idx = find_header_row(ws)

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    header = rows[header_row_idx - 1]
    header = [str(h).replace("\n", " ").strip() if h else f"col{i}" for i, h in enumerate(header)]

    data_rows = rows[header_row_idx:]

    # dedupe column names (the sheet reuses labels like "Emergency (FAE)" twice)
    seen = {}
    deduped_header = []
    for h in header:
        if h in seen:
            seen[h] += 1
            deduped_header.append(f"{h}_{seen[h]}")
        else:
            seen[h] = 0
            deduped_header.append(h)

    df = pd.DataFrame(data_rows, columns=deduped_header)
    header = deduped_header

    code_col = header[0]
    df = df[df[code_col].astype(str).str.startswith(tuple(code_prefix_filter), na=False)]
    return df


def extract_year(year_label, path):
    # 3-character sheet for the F50 total
    df3 = parse_sheet(path, "Primary Diagnosis 3 Character", ["F50"])
    # 4-character sheet for subtypes
    df4 = parse_sheet(path, "Primary Diagnosis 4 Character", ["F50."])

    combined = pd.concat([df3, df4], ignore_index=True)
    code_col = combined.columns[0]
    combined = combined.rename(columns={code_col: "Code"})
    combined["Code"] = combined["Code"].astype(str).str.strip()
    combined["diagnosis_label"] = combined["Code"].map(ED_CODE_LABELS)
    combined["financial_year"] = year_label

    # column names vary across years ("Admissions" vs "Finished Admission
    # Episodes", "Male" vs "Male (FCE)", etc.) -> match flexibly by keyword
    def find_col(cols, keywords, exclude=None):
        exclude = exclude or []
        for c in cols:
            cl = c.lower()
            if any(k in cl for k in keywords) and not any(x in cl for x in exclude):
                return c
        return None

    cols = list(combined.columns)
    fce_col = find_col(cols, ["finished consultant"])
    fae_col = find_col(cols, ["admission"])
    male_col = find_col(cols, ["male"], exclude=["female"])
    female_col = find_col(cols, ["female"])

    rename_map = {}
    if fce_col: rename_map[fce_col] = "FCE_SUM"
    if fae_col: rename_map[fae_col] = "FAE_SUM"
    if male_col: rename_map[male_col] = "FCE_Male_Sum"
    if female_col: rename_map[female_col] = "FCE_Female_Sum"
    combined = combined.rename(columns=rename_map)

    out_cols = ["financial_year", "Code", "diagnosis_label", "FCE_SUM", "FAE_SUM",
                "FCE_Male_Sum", "FCE_Female_Sum"]
    out_cols = [c for c in out_cols if c in combined.columns]
    return combined[out_cols]


def main():
    all_years = []
    for year_label, path in FILES.items():
        try:
            df = extract_year(year_label, path)
            all_years.append(df)
            print(f"Parsed {year_label}: {len(df)} rows")
        except Exception as e:
            print(f"FAILED {year_label}: {e}")

    result = pd.concat(all_years, ignore_index=True)
    result.to_csv("data/processed/nhs_ed_subtypes_multiyear.csv", index=False)
    print("\nSaved -> data/processed/nhs_ed_subtypes_multiyear.csv")
    print(result[result["Code"] == "F50"][["financial_year", "FCE_SUM", "FAE_SUM"]])


if __name__ == "__main__":
    main()
